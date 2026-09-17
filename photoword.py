import sys
import os
import json
import tempfile
import platform
import subprocess
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QListWidget, QListWidgetItem, QFileDialog,
    QMessageBox, QFrame, QSizePolicy, QProgressBar,
    QLineEdit, QCheckBox, QButtonGroup, QDialog, QScrollArea,
    QGraphicsDropShadowEffect, QAbstractItemView
)
from PySide6.QtCore import Qt, QSize, QTimer, QEvent
from PySide6.QtGui import (
    QIcon, QPixmap, QColor, QFont, QPainter, QLinearGradient, QBrush,
    QPalette, QPen
)
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageQt

# ---------- Информация ----------
APP_NAME = "Photo to Word"
APP_VERSION = "2.2.2"
APP_YEAR = "2026"
APP_AUTHOR = "Дмитрий Королев"
APP_TG = "@mr_dimakorolev"
APP_EMAIL = "mr.dimakorolev@yandex.ru"
APP_CLOUDTIPS = "https://pay.cloudtips.ru/p/56c32d56"
DEFAULT_DOC_NAME = "Photoword"

SETTINGS_FILE = os.path.join(os.path.expanduser("~"), ".photoword_settings.json")


# ---------- Локализация ----------
TRANSLATIONS = {
    "ru": {
        "app_title": "Photo to Word",
        "card1": "1. Выбор фотографий",
        "select": "📁 Выбрать фото",
        "remove": "🗑 Удалить",
        "clear": "🧹 Очистить всё",
        "about": "ℹ  О программе",
        "preview_placeholder": "Выберите фото\nв списке",
        "btn_up": "⬆ Вверх",
        "btn_down": "⬇ Вниз",
        "btn_rotate_tip": "Повернуть влево на 90°",
        "count": "Выбрано файлов: {n}",
        "card2": "2. Настройки",
        "height": "Высота (см):",
        "width": "Ширина (см):",
        "table_width": "Ширина таблицы (см):",
        "table_width_hint": "максимум 17 см",
        "table_width_hint_portrait": "максимум 17 см",
        "table_width_hint_landscape": "максимум 26 см",
        "keep_aspect": "Сохранять пропорции",
        "orientation": "Ориентация:",
        "portrait": "Книжная",
        "landscape": "Альбомная",
        "columns": "Колонок:",
        "col_1": "1 в ряд",
        "col_2": "2 в ряд",
        "col_3": "3 в ряд",
        "hint": "Доступно на странице: {av:.1f} см. На 1 колонку при {cols} в ряд: ≈ {cw:.1f} см.",
        "card3": "3. Экспорт",
        "create": "📄 Создать документ Word",
        "status_ready": "Готов к работе",
        "status_creating": "Создаю документ...",
        "status_done": "Готово! Открываю документ...",
        "status_processing": "Обработано: {i}/{n}",
        "about_title": "О программе",
        "about_author": "Автор",
        "about_tg": "Telegram",
        "about_email": "E-mail",
        "about_license": "Программа предоставляется «как есть»<br>Запрещено распространение без согласия автора",
        "about_donate": "Поддержать проект",
        "about_donate_hint": "Отсканируйте QR-код или перейдите по ссылке:",
        "about_close": "Закрыть",
        "dialog_move_title": "Перемещение",
        "dialog_move_msg": "Сначала выберите фото в списке.",
        "dialog_rotate_title": "Поворот",
        "dialog_rotate_msg": "Сначала выберите фото в списке.",
        "dialog_rotate_err": "Ошибка поворота",
        "dialog_no_files_title": "Нет файлов",
        "dialog_no_files_msg": "Сначала выберите фотографии!",
        "dialog_error_title": "Ошибка",
        "dialog_error_size": "Введите корректные положительные числа!",
        "dialog_error_table": "Ширина таблицы должна быть числом.",
        "dialog_save_title": "Сохранить документ как",
        "dialog_save_filter": "Документ Word (*.docx)",
        "dialog_open_err": "Не удалось открыть",
        "dialog_skip": "Пропуск",
        "dialog_skip_msg": "Не удалось: {name}\n{err}",
        "dialog_doc_err": "Не удалось создать документ:\n{err}",
    },
    "en": {
        "app_title": "Photo to Word",
        "card1": "1. Choose photos",
        "select": "📁 Choose photos",
        "remove": "🗑 Delete",
        "clear": "🧹 Clear all",
        "about": "ℹ  About",
        "preview_placeholder": "Select a photo\nin the list",
        "btn_up": "⬆ Up",
        "btn_down": "⬇ Down",
        "btn_rotate_tip": "Rotate 90° left",
        "count": "Files selected: {n}",
        "card2": "2. Settings",
        "height": "Height (cm):",
        "width": "Width (cm):",
        "table_width": "Table width (cm):",
        "table_width_hint": "max 17 cm",
        "table_width_hint_portrait": "max 17 cm",
        "table_width_hint_landscape": "max 26 cm",
        "keep_aspect": "Keep aspect ratio",
        "orientation": "Orientation:",
        "portrait": "Portrait",
        "landscape": "Landscape",
        "columns": "Columns:",
        "col_1": "1 per row",
        "col_2": "2 per row",
        "col_3": "3 per row",
        "hint": "Available on page: {av:.1f} cm. Per column at {cols} per row: ≈ {cw:.1f} cm.",
        "card3": "3. Export",
        "create": "📄 Create Word document",
        "status_ready": "Ready",
        "status_creating": "Creating document...",
        "status_done": "Done! Opening document...",
        "status_processing": "Processed: {i}/{n}",
        "about_title": "About",
        "about_author": "Author",
        "about_tg": "Telegram",
        "about_email": "E-mail",
        "about_license": "The program is provided \"as is\"<br>Distribution without the author's consent is prohibited",
        "about_donate": "Support the project",
        "about_donate_hint": "Scan the QR code or follow the link:",
        "about_close": "Close",
        "dialog_move_title": "Move",
        "dialog_move_msg": "Select a photo in the list first.",
        "dialog_rotate_title": "Rotate",
        "dialog_rotate_msg": "Select a photo in the list first.",
        "dialog_rotate_err": "Rotation error",
        "dialog_no_files_title": "No files",
        "dialog_no_files_msg": "Choose photos first!",
        "dialog_error_title": "Error",
        "dialog_error_size": "Enter valid positive numbers!",
        "dialog_error_table": "Table width must be a number.",
        "dialog_save_title": "Save document as",
        "dialog_save_filter": "Word Document (*.docx)",
        "dialog_open_err": "Cannot open",
        "dialog_skip": "Skipped",
        "dialog_skip_msg": "Failed: {name}\n{err}",
        "dialog_doc_err": "Failed to create document:\n{err}",
    },
}


# ---------- Настройки ----------
def load_settings() -> dict:
    try:
        if os.path.exists(SETTINGS_FILE):
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
    except Exception:
        pass
    return {}


def save_settings(data: dict):
    try:
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def resource_path(relative_path):
    if getattr(sys, "frozen", False):
        base = getattr(sys, "_MEIPASS", os.path.dirname(sys.executable))
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, relative_path)


ICON_FILE = resource_path("icon.ico")
LOGO_FILE = resource_path("logo.png")
QR_FILE = resource_path("qr.png")


# ---------- Глобальные переменные локализации/темы ----------
_settings = load_settings()
LANG = _settings.get("lang", "ru")


def tr(key: str, **kwargs) -> str:
    text = TRANSLATIONS.get(LANG, TRANSLATIONS["ru"]).get(key, key)
    if kwargs:
        try:
            text = text.format(**kwargs)
        except Exception:
            pass
    return text


# ---------- Генерация PNG-галочки ----------
def ensure_check_icon(size=32) -> str:
    path = os.path.join(tempfile.gettempdir(), "photoword_check.png")
    if os.path.exists(path):
        return path

    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    w = size
    points = [
        (int(w * 0.22), int(w * 0.52)),
        (int(w * 0.42), int(w * 0.72)),
        (int(w * 0.80), int(w * 0.28)),
    ]
    draw.line(points, fill=(255, 255, 255, 255),
              width=max(2, int(w * 0.14)), joint="curve")
    for (cx, cy) in (points[0], points[-1]):
        r = max(1, int(w * 0.07))
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=(255, 255, 255, 255))
    img.save(path, "PNG")
    return path


CHECK_ICON_PATH = ensure_check_icon()
CHECK_ICON_URL = CHECK_ICON_PATH.replace("\\", "/")


# ---------- Тема ----------
class Theme:
    def __init__(self, dark=False):
        self.dark = dark
        self.update()

    def update(self):
        if self.dark:
            self.BG_GRAD_1 = "#2E1F4A"
            self.BG_GRAD_2 = "#442A66"
            self.BG_GRAD_3 = "#5A3A7A"

            self.BG_CARD = "rgba(80, 60, 120, 0.45)"
            self.BG_CARD_BORDER = "rgba(180, 150, 230, 0.35)"

            self.BG_INPUT = "rgba(255, 255, 255, 0.08)"
            self.INPUT_BORDER = "rgba(180, 150, 230, 0.35)"

            self.TEXT = "#F0EBFF"
            self.TEXT_MUTED = "#B8A8D8"
            self.FOOTER = "#8A7AB0"

            self.ACCENT = "#A78BFA"
            self.ACCENT_HOVER = "#BCA5FF"
            self.ACCENT_PRESS = "#8E6FE8"

            self.DANGER = "#F38BA8"
            self.DANGER_HOVER = "#FFA0BC"

            self.NEUTRAL = "rgba(255, 255, 255, 0.10)"
            self.NEUTRAL_HOVER = "rgba(255, 255, 255, 0.18)"

            self.TREE_SELECT = "rgba(167, 139, 250, 0.35)"
            self.TREE_HOVER = "rgba(255, 255, 255, 0.06)"

            self.PREVIEW_BG = "rgba(0, 0, 0, 0.15)"
            self.PREVIEW_BORDER = "rgba(180, 150, 230, 0.30)"

            self.DIALOG_BG = "#2A1E44"
            self.DIALOG_TEXT = "#F0EBFF"
            self.DIALOG_BTN_BG = "#3A2B5A"
            self.DIALOG_BTN_HOVER = "#4A3A6A"
            self.DIALOG_BTN_TEXT = "#F0EBFF"
        else:
            self.BG_GRAD_1 = "#FCE9DC"
            self.BG_GRAD_2 = "#E8E4F5"
            self.BG_GRAD_3 = "#DCE9F5"

            self.BG_CARD = "rgba(255, 255, 255, 0.55)"
            self.BG_CARD_BORDER = "rgba(255, 255, 255, 0.80)"

            self.BG_INPUT = "rgba(255, 255, 255, 0.60)"
            self.INPUT_BORDER = "rgba(200, 190, 220, 0.50)"

            self.TEXT = "#3A3450"
            self.TEXT_MUTED = "#8A82A8"
            self.FOOTER = "#A8A2C0"

            self.ACCENT = "#7C6FE0"
            self.ACCENT_HOVER = "#8E82E8"
            self.ACCENT_PRESS = "#6A5CC8"

            self.DANGER = "#E85D75"
            self.DANGER_HOVER = "#F07188"

            self.NEUTRAL = "rgba(255, 255, 255, 0.50)"
            self.NEUTRAL_HOVER = "rgba(255, 255, 255, 0.70)"

            self.TREE_SELECT = "rgba(124, 111, 224, 0.20)"
            self.TREE_HOVER = "rgba(255, 255, 255, 0.40)"

            self.PREVIEW_BG = "rgba(255, 255, 255, 0.40)"
            self.PREVIEW_BORDER = "rgba(255, 255, 255, 0.70)"

            self.DIALOG_BG = "#FFFFFF"
            self.DIALOG_TEXT = "#3A3450"
            self.DIALOG_BTN_BG = "#EDEAF7"
            self.DIALOG_BTN_HOVER = "#DDD6F0"
            self.DIALOG_BTN_TEXT = "#3A3450"


theme = Theme(dark=_settings.get("dark", False))

THUMB_SIZE = 80


def font(size, weight=QFont.Normal):
    f = QFont("Segoe UI", int(size))
    f.setWeight(weight)
    return f


# ---------- Палитра приложения ----------
def apply_palette(app: QApplication):
    p = QPalette()
    bg = QColor(theme.DIALOG_BG)
    fg = QColor(theme.DIALOG_TEXT)
    base = QColor(theme.DIALOG_BG)
    btn = QColor(theme.DIALOG_BTN_BG)
    btn_text = QColor(theme.DIALOG_BTN_TEXT)
    accent = QColor(theme.ACCENT if theme.ACCENT.startswith("#") else "#7C6FE0")

    p.setColor(QPalette.Window, bg)
    p.setColor(QPalette.WindowText, fg)
    p.setColor(QPalette.Base, base)
    p.setColor(QPalette.AlternateBase, btn)
    p.setColor(QPalette.ToolTipBase, bg)
    p.setColor(QPalette.ToolTipText, fg)
    p.setColor(QPalette.Text, fg)
    p.setColor(QPalette.Button, btn)
    p.setColor(QPalette.ButtonText, btn_text)
    p.setColor(QPalette.BrightText, QColor("#FF5555"))
    p.setColor(QPalette.Link, accent)
    p.setColor(QPalette.Highlight, accent)
    p.setColor(QPalette.HighlightedText, QColor("#FFFFFF"))
    p.setColor(QPalette.PlaceholderText,
               QColor(theme.TEXT_MUTED if theme.TEXT_MUTED.startswith("#") else "#888"))
    app.setPalette(p)


# ---------- Флаги ----------
def make_ru_flag_pixmap(w=36, h=24) -> QPixmap:
    pix = QPixmap(w, h)
    pix.fill(Qt.transparent)
    painter = QPainter(pix)
    painter.setRenderHint(QPainter.Antialiasing)
    rect = pix.rect().adjusted(0, 0, -1, -1)
    painter.setPen(Qt.NoPen)
    third = h / 3
    painter.setBrush(QColor("#FFFFFF"))
    painter.drawRect(0, 0, w, int(third) + 1)
    painter.setBrush(QColor("#0039A6"))
    painter.drawRect(0, int(third), w, int(third) + 1)
    painter.setBrush(QColor("#D52B1E"))
    painter.drawRect(0, int(third * 2), w, int(third) + 1)
    painter.setPen(QPen(QColor(0, 0, 0, 60), 1))
    painter.setBrush(Qt.NoBrush)
    painter.drawRect(rect)
    painter.end()
    return pix


def make_uk_flag_pixmap(w=36, h=24) -> QPixmap:
    pix = QPixmap(w, h)
    pix.fill(Qt.transparent)
    painter = QPainter(pix)
    painter.setRenderHint(QPainter.Antialiasing)
    painter.fillRect(0, 0, w, h, QColor("#012169"))
    pen = QPen(QColor("#FFFFFF"), max(3, h // 4))
    painter.setPen(pen)
    painter.drawLine(0, 0, w, h)
    painter.drawLine(w, 0, 0, h)
    pen = QPen(QColor("#C8102E"), max(1, h // 10))
    painter.setPen(pen)
    painter.drawLine(0, 0, w, h)
    painter.drawLine(w, 0, 0, h)
    painter.setPen(Qt.NoPen)
    painter.setBrush(QColor("#FFFFFF"))
    painter.drawRect(0, h // 2 - h // 6, w, h // 3)
    painter.drawRect(w // 2 - w // 9, 0, w // 4 + w // 18, h)
    painter.setBrush(QColor("#C8102E"))
    painter.drawRect(0, h // 2 - h // 10, w, h // 5)
    painter.drawRect(w // 2 - w // 14, 0, w // 7, h)
    painter.setPen(QPen(QColor(0, 0, 0, 60), 1))
    painter.setBrush(Qt.NoBrush)
    painter.drawRect(0, 0, w - 1, h - 1)
    painter.end()
    return pix


class FlagButton(QPushButton):
    def __init__(self, lang: str, parent=None):
        super().__init__(parent)
        self.lang = lang
        self.setCursor(Qt.PointingHandCursor)
        self.setFixedSize(40, 40)
        self.setObjectName("FlagBtn")
        self._pix = None
        self.refresh()

    def refresh(self):
        if self.lang == "ru":
            self._pix = make_ru_flag_pixmap(26, 18)
        else:
            self._pix = make_uk_flag_pixmap(26, 18)
        self.setIcon(QIcon(self._pix))
        self.setIconSize(QSize(26, 18))
        self.setText("")


# ---------- Фон-градиент ----------
class GradientWidget(QWidget):
    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        grad = QLinearGradient(0, 0, self.width(), self.height())
        grad.setColorAt(0.0, QColor(theme.BG_GRAD_1))
        grad.setColorAt(0.5, QColor(theme.BG_GRAD_2))
        grad.setColorAt(1.0, QColor(theme.BG_GRAD_3))
        painter.fillRect(self.rect(), QBrush(grad))
        super().paintEvent(event)


# ---------- Кнопка ----------
class RoundedButton(QPushButton):
    def __init__(self, text, parent=None, bg=None, fg=None,
                 hover=None, press=None, radius=10, height=36,
                 width=None):
        super().__init__(text, parent)
        self.setCursor(Qt.PointingHandCursor)
        self.setFixedHeight(height)
        if width:
            self.setFixedWidth(width)
        self.radius = radius
        self._bg = bg or theme.ACCENT
        self._fg = fg or "#FFFFFF"
        self._hover = hover or theme.ACCENT_HOVER
        self._press = press or theme.ACCENT_PRESS
        self._apply_style()

    def _apply_style(self):
        self.setStyleSheet(f"""
            QPushButton {{
                background-color: {self._bg};
                color: {self._fg};
                border: none;
                border-radius: {self.radius}px;
                padding: 8px 16px;
                font-family: 'Segoe UI';
                font-size: 10pt;
                font-weight: bold;
            }}
            QPushButton:hover {{ background-color: {self._hover}; }}
            QPushButton:pressed {{ background-color: {self._press}; }}
            QPushButton:disabled {{
                background-color: {theme.NEUTRAL};
                color: {theme.TEXT_MUTED};
            }}
        """)

    def update_colors(self, bg=None, fg=None, hover=None, press=None):
        if bg: self._bg = bg
        if fg: self._fg = fg
        if hover: self._hover = hover
        if press: self._press = press
        self._apply_style()


# ---------- Карточка ----------
class Card(QFrame):
    def __init__(self, title="", parent=None):
        super().__init__(parent)
        self.setObjectName("Card")
        self.setAttribute(Qt.WA_StyledBackground, True)

        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(28)
        shadow.setColor(QColor(120, 100, 160, 60) if not theme.dark
                        else QColor(0, 0, 0, 80))
        shadow.setOffset(0, 6)
        self.setGraphicsEffect(shadow)

        self.layout = QVBoxLayout(self)
        self.layout.setContentsMargins(16, 14, 16, 14)
        self.layout.setSpacing(10)

        self.title_label = None
        if title:
            self.title_label = QLabel(title)
            self.title_label.setFont(font(12, QFont.Bold))
            self.layout.addWidget(self.title_label)

        self.update_theme()

    def set_title(self, title: str):
        if self.title_label:
            self.title_label.setText(title)

    def update_theme(self):
        self.setStyleSheet(f"""
            QFrame#Card {{
                background-color: {theme.BG_CARD};
                border-radius: 16px;
                border: 1px solid {theme.BG_CARD_BORDER};
            }}
        """)
        if self.title_label:
            self.title_label.setStyleSheet(
                f"color: {theme.TEXT}; background: transparent;")

        shadow = self.graphicsEffect()
        if shadow:
            if theme.dark:
                shadow.setColor(QColor(0, 0, 0, 80))
            else:
                shadow.setColor(QColor(120, 100, 160, 60))


# ---------- Диалог «О программе» ----------
class AboutDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle(tr("about_title"))
        self.setModal(True)
        self.setMinimumWidth(420)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 24, 24, 24)
        lay.setSpacing(12)

        if os.path.exists(LOGO_FILE):
            logo = QLabel()
            pix = QPixmap(LOGO_FILE).scaled(
                80, 80, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            logo.setPixmap(pix)
            logo.setAlignment(Qt.AlignCenter)
            lay.addWidget(logo)

        title = QLabel(APP_NAME)
        title.setAlignment(Qt.AlignCenter)
        title.setFont(font(15, QFont.Bold))
        title.setObjectName("AboutTitle")
        lay.addWidget(title)

        sub = QLabel(f"© {APP_AUTHOR}, {APP_YEAR}  •  v{APP_VERSION}")
        sub.setAlignment(Qt.AlignCenter)
        sub.setObjectName("Muted")
        lay.addWidget(sub)

        info = QLabel(
            f"<b>{tr('about_author')}:</b> {APP_AUTHOR}<br>"
            f"<b>{tr('about_tg')}:</b> {APP_TG}<br>"
            f"<b>{tr('about_email')}:</b> {APP_EMAIL}"
        )
        info.setAlignment(Qt.AlignCenter)
        info.setObjectName("AboutInfo")
        lay.addWidget(info)

        lic = QLabel(tr("about_license"))
        lic.setAlignment(Qt.AlignCenter)
        lic.setObjectName("Muted")
        lic.setWordWrap(True)
        lay.addWidget(lic)

        line = QFrame()
        line.setFrameShape(QFrame.HLine)
        line.setObjectName("AboutSep")
        lay.addWidget(line)

        donate = QLabel(tr("about_donate"))
        donate.setAlignment(Qt.AlignCenter)
        donate.setFont(font(12, QFont.Bold))
        donate.setObjectName("AboutDonate")
        lay.addWidget(donate)

        hint = QLabel(tr("about_donate_hint"))
        hint.setAlignment(Qt.AlignCenter)
        hint.setObjectName("Muted")
        hint.setWordWrap(True)
        lay.addWidget(hint)

        if os.path.exists(QR_FILE):
            qr = QLabel()
            pix = QPixmap(QR_FILE).scaled(
                180, 180, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            qr.setPixmap(pix)
            qr.setAlignment(Qt.AlignCenter)
            lay.addWidget(qr)

        link = QLabel(
            f'<a href="{APP_CLOUDTIPS}" style="color:{theme.ACCENT};'
            f' text-decoration:none;">{APP_CLOUDTIPS}</a>'
        )
        link.setAlignment(Qt.AlignCenter)
        link.setOpenExternalLinks(True)
        link.setObjectName("AboutLink")
        lay.addWidget(link)

        btn_row = QHBoxLayout()
        btn_row.addStretch()
        close_btn = RoundedButton(tr("about_close"), bg=theme.ACCENT,
                                  height=36, width=120)
        close_btn.clicked.connect(self.accept)
        btn_row.addWidget(close_btn)
        btn_row.addStretch()
        lay.addLayout(btn_row)

        self.setStyleSheet(self._qss())

    def _qss(self):
        return f"""
            QDialog {{
                background-color: {theme.DIALOG_BG};
            }}
            QLabel#AboutTitle {{ color: {theme.TEXT}; }}
            QLabel#AboutInfo  {{ color: {theme.TEXT}; }}
            QLabel#AboutDonate {{ color: {theme.ACCENT}; }}
            QLabel#AboutLink  {{ color: {theme.ACCENT}; }}
            QLabel#Muted {{ color: {theme.TEXT_MUTED}; }}
            QFrame#AboutSep {{
                background-color: {theme.INPUT_BORDER};
                max-height: 1px;
                border: none;
            }}
        """


# ---------- Главное окно ----------
class PhotoToWordApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(
            f"{APP_NAME} — v{APP_VERSION} | © {APP_AUTHOR}, {APP_YEAR}")

        if os.path.exists(ICON_FILE):
            self.setWindowIcon(QIcon(ICON_FILE))

        self.selected_files = []
        self.last_dir = ""
        self.current_preview_path = None

        self._preview_timer = QTimer(self)
        self._preview_timer.setSingleShot(True)
        self._preview_timer.setInterval(60)
        self._preview_timer.timeout.connect(self._render_preview)

        self._setup_geometry()
        self._build_ui()
        self._apply_theme()
        self._apply_language()

        self.theme_btn.setText("☀" if theme.dark else "🌙")

    # ------------------------------------------------------------------
    def _setup_geometry(self):
        """Адаптивная геометрия с учётом DPI и реального размера экрана."""
        screen = QApplication.primaryScreen()
        if screen:
            avail = screen.availableGeometry()
            aw, ah = avail.width(), avail.height()
        else:
            aw, ah = 1280, 800
            avail = None

        # Окно почти во весь экран, но с запасом.
        # avail уже в логических пикселях — при DPI 125% на 1366x768
        # он будет ~1092x614, что корректно.
        w = min(int(aw * 0.96), 1500)
        h = min(int(ah * 0.96), 1200)

        # Минимум не должен превышать экран.
        min_w = min(720, aw)
        min_h = min(520, ah)
        self.setMinimumSize(min_w, min_h)

        w = max(w, min_w)
        h = max(h, min_h)
        self.resize(w, h)

        if avail is not None:
            x = avail.x() + (aw - w) // 2
            y = avail.y() + (ah - h) // 2
            self.move(x, y)

    # ------------------------------------------------------------------
    def _build_ui(self):
        central = GradientWidget()
        self.setCentralWidget(central)

        outer = QVBoxLayout(central)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        # QScrollArea — страховка на маленьких экранах / высоком DPI.
        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.scroll.setObjectName("MainScroll")
        outer.addWidget(self.scroll)

        inner = GradientWidget()
        self.scroll.setWidget(inner)

        root = QVBoxLayout(inner)
        root.setContentsMargins(14, 10, 14, 10)
        root.setSpacing(10)

        # ===== Топбар =====
        topbar = QHBoxLayout()
        topbar.setSpacing(8)

        if os.path.exists(LOGO_FILE):
            logo_lbl = QLabel()
            pix = QPixmap(LOGO_FILE).scaled(
                56, 56, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            logo_lbl.setPixmap(pix)
            logo_lbl.setObjectName("TopLogo")
            topbar.addWidget(logo_lbl)

        topbar.addStretch()

        self.btn_about = RoundedButton(tr("about"),
                                       bg=theme.NEUTRAL, fg=theme.TEXT,
                                       hover=theme.NEUTRAL_HOVER, height=36)
        self.btn_about.clicked.connect(self.show_about)
        topbar.addWidget(self.btn_about)

        self.btn_lang = FlagButton(LANG)
        self.btn_lang.clicked.connect(self.toggle_language)
        topbar.addWidget(self.btn_lang)

        self.theme_btn = QPushButton("🌙")
        self.theme_btn.setFixedSize(40, 40)
        self.theme_btn.setCursor(Qt.PointingHandCursor)
        self.theme_btn.setObjectName("ThemeBtn")
        self.theme_btn.clicked.connect(self.toggle_theme)
        topbar.addWidget(self.theme_btn)

        root.addLayout(topbar)

        # ===== Карточка 1 =====
        self.card1 = Card(tr("card1"))
        self.card1.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.card1.setMinimumHeight(240)

        btn_row = QHBoxLayout()
        btn_row.setSpacing(8)
        self.btn_select = RoundedButton(tr("select"), bg=theme.ACCENT)
        self.btn_select.clicked.connect(self.select_files)
        btn_row.addWidget(self.btn_select)

        self.btn_remove = RoundedButton(tr("remove"), bg=theme.DANGER,
                                        hover=theme.DANGER_HOVER)
        self.btn_remove.clicked.connect(self.remove_selected)
        btn_row.addWidget(self.btn_remove)

        self.btn_clear = RoundedButton(tr("clear"), bg=theme.NEUTRAL,
                                       fg=theme.TEXT, hover=theme.NEUTRAL_HOVER)
        self.btn_clear.clicked.connect(self.clear_all)
        btn_row.addWidget(self.btn_clear)
        btn_row.addStretch()
        self.card1.layout.addLayout(btn_row)

        content_row = QHBoxLayout()
        content_row.setSpacing(10)

        self.list_widget = QListWidget()
        self.list_widget.setIconSize(QSize(THUMB_SIZE, THUMB_SIZE))
        self.list_widget.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.list_widget.setObjectName("PhotoList")
        self.list_widget.itemSelectionChanged.connect(self.on_select)
        self.list_widget.setMinimumWidth(150)
        self.list_widget.setMinimumHeight(140)
        self.list_widget.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Expanding)
        content_row.addWidget(self.list_widget, 1)

        right_col = QVBoxLayout()
        right_col.setSpacing(8)

        preview_frame = QFrame()
        preview_frame.setObjectName("PreviewFrame")
        preview_frame.setMinimumHeight(160)
        preview_layout = QVBoxLayout(preview_frame)
        preview_layout.setContentsMargins(8, 8, 8, 8)
        preview_layout.setSpacing(6)

        self.preview_label = QLabel(tr("preview_placeholder"))
        self.preview_label.setAlignment(Qt.AlignCenter)
        self.preview_label.setObjectName("PreviewLabel")
        self.preview_label.setMinimumSize(180, 130)
        self.preview_label.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.preview_label.installEventFilter(self)
        preview_layout.addWidget(self.preview_label, 1)

        self.preview_info = QLabel("")
        self.preview_info.setObjectName("PreviewInfo")
        self.preview_info.setAlignment(Qt.AlignCenter)
        self.preview_info.setWordWrap(True)
        preview_layout.addWidget(self.preview_info)

        right_col.addWidget(preview_frame, 1)

        ctrl_row = QHBoxLayout()
        ctrl_row.setSpacing(6)
        ctrl_row.addStretch()

        self.btn_up = RoundedButton(tr("btn_up"), bg=theme.NEUTRAL, fg=theme.TEXT,
                                    hover=theme.NEUTRAL_HOVER, height=34,
                                    width=100)
        self.btn_up.clicked.connect(lambda: self.move_selected(-1))
        ctrl_row.addWidget(self.btn_up)

        self.btn_down = RoundedButton(tr("btn_down"), bg=theme.NEUTRAL, fg=theme.TEXT,
                                      hover=theme.NEUTRAL_HOVER, height=34,
                                      width=100)
        self.btn_down.clicked.connect(lambda: self.move_selected(1))
        ctrl_row.addWidget(self.btn_down)

        self.btn_rotate = QPushButton("↺")
        self.btn_rotate.setFixedSize(34, 34)
        self.btn_rotate.setCursor(Qt.PointingHandCursor)
        self.btn_rotate.setObjectName("RotateBtn")
        self.btn_rotate.setToolTip(tr("btn_rotate_tip"))
        self.btn_rotate.clicked.connect(self.rotate_selected_left)
        ctrl_row.addWidget(self.btn_rotate)

        ctrl_row.addStretch()
        right_col.addLayout(ctrl_row)

        content_row.addLayout(right_col, 2)
        self.card1.layout.addLayout(content_row, 1)

        self.count_label = QLabel(tr("count", n=0))
        self.count_label.setObjectName("CountLabel")
        self.card1.layout.addWidget(self.count_label)

        root.addWidget(self.card1, 1)

        # ===== Карточка 2 =====
        self.card2 = Card(tr("card2"))
        self.card2.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)

        row1 = QHBoxLayout()
        row1.setSpacing(6)
        self.lbl_height = QLabel(tr("height"))
        self.lbl_height.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row1.addWidget(self.lbl_height)
        self.height_input = QLineEdit("6")
        self.height_input.setMinimumWidth(56)
        self.height_input.setMaximumWidth(90)
        self.height_input.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        row1.addWidget(self.height_input)
        row1.addSpacing(6)
        self.lbl_width = QLabel(tr("width"))
        self.lbl_width.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row1.addWidget(self.lbl_width)
        self.width_input = QLineEdit("8")
        self.width_input.setMinimumWidth(56)
        self.width_input.setMaximumWidth(90)
        self.width_input.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        row1.addWidget(self.width_input)
        row1.addStretch()
        self.card2.layout.addLayout(row1)

        row2 = QHBoxLayout()
        row2.setSpacing(6)
        self.lbl_table_w = QLabel(tr("table_width"))
        self.lbl_table_w.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row2.addWidget(self.lbl_table_w)
        self.table_width_input = QLineEdit("17")
        self.table_width_input.setMinimumWidth(56)
        self.table_width_input.setMaximumWidth(90)
        self.table_width_input.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        row2.addWidget(self.table_width_input)
        self.lbl_table_w_hint = QLabel(tr("table_width_hint_portrait"))
        self.lbl_table_w_hint.setObjectName("Muted")
        self.lbl_table_w_hint.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row2.addWidget(self.lbl_table_w_hint)
        row2.addStretch()
        self.card2.layout.addLayout(row2)

        row3 = QHBoxLayout()
        self.keep_aspect = QCheckBox(tr("keep_aspect"))
        self.keep_aspect.setChecked(True)
        self.keep_aspect.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row3.addWidget(self.keep_aspect)
        row3.addStretch()
        self.card2.layout.addLayout(row3)

        row4 = QHBoxLayout()
        row4.setSpacing(8)
        self.lbl_orient = QLabel(tr("orientation"))
        self.lbl_orient.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row4.addWidget(self.lbl_orient)
        self.orient_group = QButtonGroup(self)
        self.orient_group.setExclusive(True)
        self.rb_portrait = QCheckBox(tr("portrait"))
        self.rb_portrait.setChecked(True)
        self.rb_portrait.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        self.rb_landscape = QCheckBox(tr("landscape"))
        self.rb_landscape.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        self.orient_group.addButton(self.rb_portrait)
        self.orient_group.addButton(self.rb_landscape)
        row4.addWidget(self.rb_portrait)
        row4.addWidget(self.rb_landscape)
        row4.addStretch()
        self.card2.layout.addLayout(row4)

        row5 = QHBoxLayout()
        row5.setSpacing(8)
        self.lbl_cols = QLabel(tr("columns"))
        self.lbl_cols.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
        row5.addWidget(self.lbl_cols)
        self.col_group = QButtonGroup(self)
        self.col_group.setExclusive(True)
        self.col_buttons = []
        for c, key in zip((1, 2, 3), ("col_1", "col_2", "col_3")):
            cb = QCheckBox(tr(key))
            cb.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Preferred)
            if c == 2:
                cb.setChecked(True)
            self.col_group.addButton(cb, c)
            row5.addWidget(cb)
            self.col_buttons.append(cb)
        row5.addStretch()
        self.card2.layout.addLayout(row5)

        self.hint_label = QLabel("")
        self.hint_label.setObjectName("Muted")
        self.hint_label.setWordWrap(True)
        self.card2.layout.addWidget(self.hint_label)

        self.rb_portrait.toggled.connect(self.update_hint)
        self.col_group.buttonClicked.connect(self.update_hint)

        root.addWidget(self.card2, 0)

        # ===== Карточка 3 =====
        self.card3 = Card(tr("card3"))
        self.card3.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)

        self.create_btn = RoundedButton(tr("create"),
                                        bg=theme.ACCENT, height=44)
        self.create_btn.setFont(font(12, QFont.Bold))
        self.create_btn.clicked.connect(self.create_word_doc)
        self.card3.layout.addWidget(self.create_btn, alignment=Qt.AlignCenter)

        self.progress = QProgressBar()
        self.progress.setFixedHeight(8)
        self.progress.setTextVisible(False)
        self.card3.layout.addWidget(self.progress)

        self.status_label = QLabel(tr("status_ready"))
        self.status_label.setObjectName("Muted")
        self.status_label.setAlignment(Qt.AlignCenter)
        self.status_label.setWordWrap(True)
        self.card3.layout.addWidget(self.status_label)

        footer = QLabel(f"© {APP_YEAR} • v{APP_VERSION}")
        footer.setObjectName("Footer")
        footer.setAlignment(Qt.AlignCenter)
        self.card3.layout.addWidget(footer)

        root.addWidget(self.card3, 0)

        self.update_hint()
        self.update_count()

    # ==============================================================
    def _apply_language(self):
        self.card1.set_title(tr("card1"))
        self.btn_select.setText(tr("select"))
        self.btn_remove.setText(tr("remove"))
        self.btn_clear.setText(tr("clear"))
        self.btn_about.setText(tr("about"))

        if not self.current_preview_path:
            self.preview_label.setText(tr("preview_placeholder"))

        self.btn_up.setText(tr("btn_up"))
        self.btn_down.setText(tr("btn_down"))
        self.btn_rotate.setToolTip(tr("btn_rotate_tip"))

        self.update_count()

        self.card2.set_title(tr("card2"))
        self.lbl_height.setText(tr("height"))
        self.lbl_width.setText(tr("width"))
        self.lbl_table_w.setText(tr("table_width"))
        self.keep_aspect.setText(tr("keep_aspect"))
        self.lbl_orient.setText(tr("orientation"))
        self.rb_portrait.setText(tr("portrait"))
        self.rb_landscape.setText(tr("landscape"))
        self.lbl_cols.setText(tr("columns"))
        for cb, key in zip(self.col_buttons, ("col_1", "col_2", "col_3")):
            cb.setText(tr(key))

        self.card3.set_title(tr("card3"))
        self.create_btn.setText(tr("create"))
        if self.status_label.text() in (
            TRANSLATIONS["ru"]["status_ready"], TRANSLATIONS["en"]["status_ready"]
        ):
            self.status_label.setText(tr("status_ready"))

        self.update_hint()

    def toggle_language(self):
        global LANG
        LANG = "en" if LANG == "ru" else "ru"
        _settings["lang"] = LANG
        save_settings(_settings)
        self.btn_lang.lang = LANG
        self.btn_lang.refresh()
        self._apply_language()

    def toggle_theme(self):
        theme.dark = not theme.dark
        theme.update()
        self.theme_btn.setText("☀" if theme.dark else "🌙")
        self._apply_theme()
        self.centralWidget().update()

        app = QApplication.instance()
        if app:
            apply_palette(app)
            app.setStyleSheet(self._dialog_qss())

        _settings["dark"] = theme.dark
        save_settings(_settings)

    def closeEvent(self, event):
        _settings["dark"] = theme.dark
        _settings["lang"] = LANG
        save_settings(_settings)
        super().closeEvent(event)

    def _dialog_qss(self) -> str:
        return f"""
            QMessageBox {{
                background-color: {theme.DIALOG_BG};
            }}
            QMessageBox QLabel {{
                color: {theme.DIALOG_TEXT};
                background: transparent;
                font-family: 'Segoe UI';
                font-size: 10pt;
            }}
            QMessageBox QPushButton {{
                background-color: {theme.DIALOG_BTN_BG};
                color: {theme.DIALOG_BTN_TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 6px;
                padding: 6px 18px;
                min-width: 80px;
                font-family: 'Segoe UI';
                font-size: 10pt;
                font-weight: bold;
            }}
            QMessageBox QPushButton:hover {{
                background-color: {theme.DIALOG_BTN_HOVER};
            }}
            QMessageBox QPushButton:pressed {{
                background-color: {theme.ACCENT};
                color: #FFFFFF;
            }}

            QDialog {{
                background-color: {theme.DIALOG_BG};
                color: {theme.DIALOG_TEXT};
            }}
            QDialog QLabel {{
                color: {theme.DIALOG_TEXT};
                background: transparent;
            }}

            QFileDialog {{
                background-color: {theme.DIALOG_BG};
                color: {theme.DIALOG_TEXT};
            }}
            QFileDialog QLabel {{
                color: {theme.DIALOG_TEXT};
                background: transparent;
            }}
            QFileDialog QLineEdit,
            QFileDialog QComboBox {{
                background-color: {theme.BG_INPUT};
                color: {theme.DIALOG_TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 6px;
                padding: 4px 8px;
            }}
            QFileDialog QTreeView,
            QFileDialog QListView {{
                background-color: {theme.DIALOG_BG};
                color: {theme.DIALOG_TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 6px;
                outline: none;
            }}
            QFileDialog QTreeView::item,
            QFileDialog QListView::item {{
                color: {theme.DIALOG_TEXT};
                padding: 4px;
            }}
            QFileDialog QTreeView::item:selected,
            QFileDialog QListView::item:selected {{
                background-color: {theme.ACCENT};
                color: #FFFFFF;
            }}
            QFileDialog QTreeView::item:hover,
            QFileDialog QListView::item:hover {{
                background-color: {theme.DIALOG_BTN_HOVER};
            }}
            QFileDialog QHeaderView::section {{
                background-color: {theme.DIALOG_BTN_BG};
                color: {theme.DIALOG_TEXT};
                border: none;
                padding: 4px 8px;
            }}
            QFileDialog QPushButton {{
                background-color: {theme.DIALOG_BTN_BG};
                color: {theme.DIALOG_BTN_TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 6px;
                padding: 6px 14px;
                min-width: 70px;
            }}
            QFileDialog QPushButton:hover {{
                background-color: {theme.DIALOG_BTN_HOVER};
            }}
        """

    def _apply_theme(self):
        check_url = CHECK_ICON_URL
        if not os.path.exists(CHECK_ICON_PATH):
            check_url = ensure_check_icon().replace("\\", "/")

        self.setStyleSheet(f"""
            QMainWindow {{ background-color: {theme.BG_GRAD_1}; }}

            QScrollArea#MainScroll {{
                background: transparent;
                border: none;
            }}
            QScrollArea#MainScroll > QWidget > QWidget {{
                background: transparent;
            }}

            QWidget {{
                background: transparent;
                color: {theme.TEXT};
                font-family: 'Segoe UI';
                font-size: 10pt;
            }}

            QLabel {{ background: transparent; }}
            QLabel#TopLogo {{ background: transparent; }}
            QLabel#Muted {{ color: {theme.TEXT_MUTED}; }}
            QLabel#Footer {{ color: {theme.FOOTER}; font-size: 8pt; }}
            QLabel#CountLabel {{ color: {theme.TEXT_MUTED}; }}
            QLabel#PreviewLabel {{
                color: {theme.TEXT_MUTED};
                background-color: {theme.PREVIEW_BG};
                border-radius: 10px;
                padding: 6px;
            }}
            QLabel#PreviewInfo {{ color: {theme.TEXT_MUTED}; font-size: 9pt; }}

            QPushButton#ThemeBtn, QPushButton#FlagBtn {{
                background-color: {theme.BG_CARD};
                color: {theme.TEXT};
                border: 1px solid {theme.BG_CARD_BORDER};
                border-radius: 20px;
                font-size: 14pt;
            }}
            QPushButton#ThemeBtn:hover, QPushButton#FlagBtn:hover {{
                background-color: {theme.NEUTRAL_HOVER};
            }}

            QPushButton#RotateBtn {{
                background-color: {theme.NEUTRAL};
                color: {theme.TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 17px;
                font-size: 15pt;
                font-weight: bold;
                padding: 0;
            }}
            QPushButton#RotateBtn:hover {{
                background-color: {theme.NEUTRAL_HOVER};
                border: 1px solid {theme.ACCENT};
                color: {theme.ACCENT};
            }}
            QPushButton#RotateBtn:pressed {{
                background-color: {theme.ACCENT};
                color: #FFFFFF;
            }}

            QLineEdit {{
                background-color: {theme.BG_INPUT};
                color: {theme.TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 8px;
                padding: 5px 8px;
            }}
            QLineEdit:focus {{ border: 1px solid {theme.ACCENT}; }}

            QCheckBox {{
                color: {theme.TEXT};
                background: transparent;
                spacing: 6px;
            }}
            QCheckBox::indicator {{
                width: 18px; height: 18px;
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 5px;
                background-color: {theme.BG_INPUT};
            }}
            QCheckBox::indicator:hover {{
                border: 1px solid {theme.ACCENT};
            }}
            QCheckBox::indicator:checked {{
                background-color: {theme.ACCENT};
                border: 1px solid {theme.ACCENT};
                image: url("{check_url}");
            }}
            QCheckBox::indicator:checked:hover {{
                background-color: {theme.ACCENT_HOVER};
                border: 1px solid {theme.ACCENT_HOVER};
            }}

            QListWidget#PhotoList {{
                background-color: {theme.BG_INPUT};
                color: {theme.TEXT};
                border: 1px solid {theme.INPUT_BORDER};
                border-radius: 10px;
                padding: 4px;
                outline: none;
            }}
            QListWidget#PhotoList::item {{
                padding: 3px;
                border-radius: 8px;
            }}
            QListWidget#PhotoList::item:selected {{
                background-color: {theme.TREE_SELECT};
                color: {theme.TEXT};
            }}
            QListWidget#PhotoList::item:hover {{
                background-color: {theme.TREE_HOVER};
            }}

            QFrame#PreviewFrame {{
                background-color: {theme.PREVIEW_BG};
                border: 1px solid {theme.PREVIEW_BORDER};
                border-radius: 12px;
            }}

            QProgressBar {{
                background-color: {theme.BG_INPUT};
                border: none;
                border-radius: 4px;
            }}
            QProgressBar::chunk {{
                background-color: {theme.ACCENT};
                border-radius: 4px;
            }}

            QScrollBar:vertical {{
                background: transparent;
                width: 10px;
                margin: 0;
            }}
            QScrollBar::handle:vertical {{
                background: {theme.INPUT_BORDER};
                border-radius: 5px;
                min-height: 30px;
            }}
            QScrollBar::handle:vertical:hover {{
                background: {theme.ACCENT};
            }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        self.btn_select.update_colors(bg=theme.ACCENT, hover=theme.ACCENT_HOVER,
                                      press=theme.ACCENT_PRESS)
        self.btn_remove.update_colors(bg=theme.DANGER, hover=theme.DANGER_HOVER)
        self.btn_clear.update_colors(bg=theme.NEUTRAL, fg=theme.TEXT,
                                     hover=theme.NEUTRAL_HOVER)
        self.btn_about.update_colors(bg=theme.NEUTRAL, fg=theme.TEXT,
                                     hover=theme.NEUTRAL_HOVER)
        self.btn_up.update_colors(bg=theme.NEUTRAL, fg=theme.TEXT,
                                  hover=theme.NEUTRAL_HOVER)
        self.btn_down.update_colors(bg=theme.NEUTRAL, fg=theme.TEXT,
                                    hover=theme.NEUTRAL_HOVER)
        self.create_btn.update_colors(bg=theme.ACCENT, hover=theme.ACCENT_HOVER,
                                      press=theme.ACCENT_PRESS)

        for card in self.findChildren(Card):
            card.update_theme()

    # ==============================================================
    def eventFilter(self, obj, event):
        if obj is self.preview_label and event.type() == QEvent.Resize:
            self._preview_timer.start()
        return super().eventFilter(obj, event)

    def refresh_preview(self):
        self._preview_timer.start()

    def update_hint(self):
        is_landscape = self.rb_landscape.isChecked()
        page_width = 29.7 if is_landscape else 21.0
        available = page_width - 4.0
        cols = self.col_group.checkedId()
        if cols < 1:
            cols = 2
        cell_w = available / cols
        self.hint_label.setText(tr("hint", av=available, cols=cols, cw=cell_w))

        if is_landscape:
            self.lbl_table_w_hint.setText(tr("table_width_hint_landscape"))
        else:
            self.lbl_table_w_hint.setText(tr("table_width_hint_portrait"))

    def update_count(self):
        self.count_label.setText(tr("count", n=len(self.selected_files)))

    def _make_thumbnail(self, path):
        try:
            img = Image.open(path)
            img.thumbnail((THUMB_SIZE, THUMB_SIZE), Image.LANCZOS)
            return QPixmap.fromImage(ImageQt.ImageQt(img))
        except Exception:
            return None

    def _add_file_to_list(self, path):
        if path in self.selected_files:
            return
        self.selected_files.append(path)
        item = QListWidgetItem(os.path.basename(path))
        thumb = self._make_thumbnail(path)
        if thumb:
            item.setIcon(QIcon(thumb))
        item.setToolTip(path)
        self.list_widget.addItem(item)
        self.last_dir = os.path.dirname(path)

    def select_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, "Выберите фотографии" if LANG == "ru" else "Choose photos",
            self.last_dir or "",
            "Images (*.jpg *.jpeg *.png *.bmp *.gif)")
        for f in files:
            self._add_file_to_list(f)
        self.update_count()

    def remove_selected(self):
        indices = sorted([self.list_widget.row(i)
                          for i in self.list_widget.selectedItems()],
                         reverse=True)
        for idx in indices:
            self.list_widget.takeItem(idx)
            del self.selected_files[idx]
        self.update_count()
        self.current_preview_path = None
        self.preview_label.setPixmap(QPixmap())
        self.preview_label.setText(tr("preview_placeholder"))
        self.preview_info.setText("")

    def clear_all(self):
        self.list_widget.clear()
        self.selected_files.clear()
        self.current_preview_path = None
        self.update_count()
        self.preview_label.setPixmap(QPixmap())
        self.preview_label.setText(tr("preview_placeholder"))
        self.preview_info.setText("")

    # ---------- Перемещение ----------
    def move_selected(self, direction):
        items = self.list_widget.selectedItems()
        if not items:
            QMessageBox.information(self, tr("dialog_move_title"),
                                    tr("dialog_move_msg"))
            return

        indices = sorted([self.list_widget.row(i) for i in items])

        if direction == -1 and indices[0] == 0:
            return
        if direction == 1 and indices[-1] == self.list_widget.count() - 1:
            return

        if direction == -1:
            for idx in indices:
                self._swap_items(idx, idx - 1)
            new_indices = [i - 1 for i in indices]
        else:
            for idx in reversed(indices):
                self._swap_items(idx, idx + 1)
            new_indices = [i + 1 for i in indices]

        self.list_widget.clearSelection()
        for idx in new_indices:
            item = self.list_widget.item(idx)
            if item:
                item.setSelected(True)

        first = self.list_widget.item(new_indices[0])
        if first:
            self.list_widget.scrollToItem(first, QAbstractItemView.EnsureVisible)

        self.on_select()

    def _swap_items(self, i, j):
        if i == j:
            return

        self.selected_files[i], self.selected_files[j] = \
            self.selected_files[j], self.selected_files[i]

        item_i = self.list_widget.item(i)
        item_j = self.list_widget.item(j)

        self.list_widget.takeItem(i)
        self.list_widget.takeItem(j - 1 if j > i else j)

        if i < j:
            self.list_widget.insertItem(i, item_j)
            self.list_widget.insertItem(j, item_i)
        else:
            self.list_widget.insertItem(j, item_i)
            self.list_widget.insertItem(i, item_j)

    # ---------- Поворот ----------
    def rotate_selected_left(self):
        items = self.list_widget.selectedItems()
        if not items:
            QMessageBox.information(self, tr("dialog_rotate_title"),
                                    tr("dialog_rotate_msg"))
            return

        indices = sorted([self.list_widget.row(i) for i in items])
        for idx in indices:
            path = self.selected_files[idx]
            try:
                img = Image.open(path)
                exif = img.info.get("exif", None)
                rotated_img = img.rotate(90, expand=True)
                save_kwargs = {}
                ext = os.path.splitext(path)[1].lower()
                if ext in (".jpg", ".jpeg"):
                    save_kwargs["quality"] = 95
                    if exif:
                        save_kwargs["exif"] = exif
                rotated_img.save(path, **save_kwargs)
            except Exception as e:
                QMessageBox.warning(self, tr("dialog_rotate_err"),
                                    f"{os.path.basename(path)}:\n{e}")

        for idx in indices:
            path = self.selected_files[idx]
            item = self.list_widget.item(idx)
            thumb = self._make_thumbnail(path)
            if thumb and item:
                item.setIcon(QIcon(thumb))

        items = self.list_widget.selectedItems()
        if items:
            idx = self.list_widget.row(items[0])
            self.current_preview_path = self.selected_files[idx]
            self._preview_timer.start()

    # ---------- Превью ----------
    def on_select(self):
        items = self.list_widget.selectedItems()
        if not items:
            return
        idx = self.list_widget.row(items[0])
        self.current_preview_path = self.selected_files[idx]
        self._preview_timer.start()

    def _render_preview(self):
        path = self.current_preview_path
        if not path or not os.path.exists(path):
            return
        try:
            sz = self.preview_label.size()
            box_w = max(sz.width() - 12, 80)
            box_h = max(sz.height() - 12, 80)

            img = Image.open(path)
            w, h = img.size
            img.thumbnail((box_w, box_h), Image.LANCZOS)
            pix = QPixmap.fromImage(ImageQt.ImageQt(img))
            self.preview_label.setPixmap(pix)
            self.preview_label.setText("")
            size_kb = os.path.getsize(path) / 1024
            self.preview_info.setText(
                f"{os.path.basename(path)} — {w}×{h} px, {size_kb:.0f} KB")
        except Exception as e:
            self.preview_label.setText(f"Error:\n{e}")
            self.preview_info.setText("")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.list_widget.selectedItems():
            self._preview_timer.start()

    # ==============================================================
    def show_about(self):
        dlg = AboutDialog(self)
        dlg.exec()

    # ==============================================================
    def create_word_doc(self):
        if not self.selected_files:
            QMessageBox.warning(self, tr("dialog_no_files_title"),
                                tr("dialog_no_files_msg"))
            return

        try:
            height_cm = float(self.height_input.text())
            width_cm = float(self.width_input.text())
            if height_cm <= 0 or width_cm <= 0:
                raise ValueError
        except ValueError:
            QMessageBox.critical(self, tr("dialog_error_title"),
                                 tr("dialog_error_size"))
            return

        cols = self.col_group.checkedId()
        if cols < 1:
            cols = 2

        table_width_str = self.table_width_input.text().strip()
        custom_table_width = None
        if table_width_str:
            try:
                custom_table_width = float(table_width_str)
            except ValueError:
                QMessageBox.critical(self, tr("dialog_error_title"),
                                     tr("dialog_error_table"))
                return

        initial_dir = self.last_dir or (
            os.path.dirname(self.selected_files[0]) if self.selected_files else "")
        default_name = DEFAULT_DOC_NAME
        if initial_dir:
            candidate = os.path.join(initial_dir, default_name + ".docx")
            counter = 1
            while os.path.exists(candidate):
                candidate = os.path.join(initial_dir,
                                         f"{default_name}_{counter}.docx")
                counter += 1
            default_name = os.path.splitext(os.path.basename(candidate))[0]

        save_path, _ = QFileDialog.getSaveFileName(
            self, tr("dialog_save_title"),
            os.path.join(initial_dir, default_name) if initial_dir else default_name,
            tr("dialog_save_filter"))
        if not save_path:
            return

        self.last_dir = os.path.dirname(save_path)
        self.create_btn.setEnabled(False)
        self.status_label.setText(tr("status_creating"))
        QApplication.processEvents()

        try:
            doc = Document()
            doc.core_properties.author = APP_AUTHOR
            doc.core_properties.last_modified_by = APP_AUTHOR
            doc.core_properties.comments = f"Created in {APP_NAME} v{APP_VERSION}"
            doc.core_properties.title = "Photos"

            section = doc.sections[0]
            if self.rb_landscape.isChecked():
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width, section.page_height = \
                    section.page_height, section.page_width
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)

            total = len(self.selected_files)
            rows_with_photos = (total + cols - 1) // cols
            total_rows = 2 * rows_with_photos - 1 if rows_with_photos > 0 else 0

            table = doc.add_table(rows=total_rows, cols=cols)
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            page_width_cm = 29.7 if self.rb_landscape.isChecked() else 21.0
            available = page_width_cm - 4.0
            table_width = min(custom_table_width, available) \
                if custom_table_width else available
            col_width = Cm(table_width / cols)

            for col in table.columns:
                col.width = col_width
            for row in table.rows:
                for cell in row.cells:
                    cell.width = col_width

            tblPr = table._tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr')
                table._tbl.insert(0, tblPr)
            borders = OxmlElement('w:tblBorders')
            for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                b = OxmlElement(f'w:{name}')
                b.set(qn('w:val'), 'none')
                b.set(qn('w:sz'), '0')
                b.set(qn('w:space'), '0')
                b.set(qn('w:color'), 'auto')
                borders.append(b)
            tblPr.append(borders)

            self.progress.setMaximum(total)

            for i, file_path in enumerate(self.selected_files):
                try:
                    img = Image.open(file_path)
                    img.verify()
                    img = Image.open(file_path)

                    photo_row = (i // cols) * 2
                    col = i % cols
                    cell = table.cell(photo_row, col)
                    for p in cell.paragraphs:
                        p.clear()
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    run = p.add_run()

                    if self.keep_aspect.isChecked():
                        iw, ih = img.size
                        ratio = min(width_cm / iw, height_cm / ih)
                        run.add_picture(file_path,
                                        width=Cm(iw * ratio),
                                        height=Cm(ih * ratio))
                    else:
                        run.add_picture(file_path,
                                        width=Cm(width_cm),
                                        height=Cm(height_cm))

                    tc_pr = cell._tc.get_or_add_tcPr()
                    v = tc_pr.makeelement(qn('w:vAlign'), {qn('w:val'): 'center'})
                    tc_pr.append(v)

                    self.progress.setValue(i + 1)
                    self.status_label.setText(
                        tr("status_processing", i=i + 1, n=total))
                    QApplication.processEvents()
                except Exception as e:
                    QMessageBox.warning(self, tr("dialog_skip"),
                                        tr("dialog_skip_msg",
                                           name=os.path.basename(file_path),
                                           err=e))
                    continue

            for empty_row in range(1, total_rows, 2):
                row = table.rows[empty_row]
                row.height = Pt(5)
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                for cell in row.cells:
                    for p in cell.paragraphs[1:]:
                        p._element.getparent().remove(p._element)
                    p = cell.paragraphs[0]
                    p.clear()
                    r = p.add_run()
                    r.font.size = Pt(1)
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

            doc.save(save_path)
            self.status_label.setText(tr("status_done"))
            self._open_in_word(save_path)
        except Exception as e:
            QMessageBox.critical(self, tr("dialog_error_title"),
                                 tr("dialog_doc_err", err=e))
        finally:
            self.create_btn.setEnabled(True)
            self.progress.setValue(0)
            self.status_label.setText(tr("status_ready"))

    @staticmethod
    def _open_in_word(path):
        try:
            if platform.system() == "Windows":
                os.startfile(path)
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", path])
            else:
                subprocess.Popen(["xdg-open", path])
        except Exception as e:
            QMessageBox.warning(None, tr("dialog_open_err"), f"{path}\n\n{e}")


# ==============================================================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    apply_palette(app)

    window = PhotoToWordApp()
    window.show()

    app.setStyleSheet(window._dialog_qss())

    QTimer.singleShot(0, window.refresh_preview)

    sys.exit(app.exec())
